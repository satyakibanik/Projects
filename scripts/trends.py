"""Compact daily aggregates from authoritative operating-day workbooks.

No hourly MW sums or inferred load shedding; missing fields stay null.
"""
import datetime as dt
import math
from water_level import extract_water_level, water_level_summary
from operational_events import extract_event_log


def number(value):
    return round(value, 6) if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value >= 0 else None


def temperature(value):
    """A signed Celsius reading; missing and non-finite values remain unavailable."""
    return round(value, 6) if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None


def daily_summary(report):
    summary = report.get('summary') or {}
    curve = report.get('en_curve') or {}
    totals = curve.get('totals') or {} if curve.get('date') == report['date'] else {}
    requirement = number(totals.get('total_requirement_mwh'))
    basis = 'En-Curve published requirement' if requirement is not None else None
    generated, unserved = number(summary.get('energy_mwh')), number(summary.get('unserved_mwh'))
    if requirement is None and generated is not None and unserved is not None:
        requirement = round(generated + unserved, 6)
        basis = 'P1 generation + reported energy not served'
    shedding = [number(row.get('shed_mw')) for row in report.get('hourly', [])]
    shedding = [value for value in shedding if value is not None]
    return {
        'date': report['date'], 'report_id': report.get('report_id'),
        'source_url': report.get('source_url'),
        'quality_notes': report.get('quality_notes', []),
        'generation_mwh': number(summary.get('energy_mwh')),
        'temperature_c': temperature(summary.get('temperature_c')),
        'water_level': water_level_summary(report),
        'demand_mwh': requirement, 'demand_basis': basis,
        'shed_peak_mw': max(shedding) if shedding else None,
        'regional_shed_mw': {region['name']: number(region.get('shed_mw')) for region in report.get('regions', [])},
        'unserved_mwh': number(summary.get('unserved_mwh')),
        'imports_mwh': number(summary.get('import_mwh')),
        'fuels': {key: number(value) for key, value in (report.get('fuels') or {}).items()},
        'routes': {route['name']: number(route.get('energy_mwh')) for route in report.get('imports', [])},
    }


def build_trends(reports, generated_at):
    # Same operating-day precedence as the archive: next-day publication, then newest ID.
    selected = {}
    for report in reports:
        day = dt.date.fromisoformat(report['date'])
        published = dt.date.fromisoformat(report.get('published_date') or report['date'])
        score = (abs((published - day).days - 1), -int(report.get('report_id') or 0))
        if day not in selected or score < selected[day][0]:
            selected[day] = (score, report)
    if not selected:
        return {'generated_at': generated_at, 'end_date': None, 'days': []}
    end = max(selected)
    start = end - dt.timedelta(days=364)
    return {'generated_at': generated_at, 'end_date': end.isoformat(),
            'days': [daily_summary(selected[day][1]) for day in sorted(selected) if day >= start]}


def extract_trend_workbook(content, item):
    """Read national totals independently of unrelated plant/region reconciliation.

    Keep raw reported values and disclose inconsistent subtotals; do not repair them.
    Only the explicitly supported P1/P4 layout is accepted.
    """
    import io
    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    if not {'P1', 'P4'} <= set(workbook.sheetnames):
        raise ValueError('Unsupported national-trends workbook layout')
    sheet = workbook['P1']
    raw_date = sheet['D6'].value
    if isinstance(raw_date, dt.datetime):
        day = raw_date.date()
    elif isinstance(raw_date, dt.date):
        day = raw_date
    else:
        day = dt.datetime.strptime(str(raw_date).strip(), '%d-%m-%Y').date()
    def energy(value):
        value = number(value)
        return round(value * 1000, 6) if value is not None else None
    fuels = {fuel:energy(sheet.cell(26, col).value) for col, fuel in enumerate(
        ['Gas','Coal','HFO','HSD','Hydro','Solar','Imports','Wind'], 5)}
    hourly = [{'time':str(row[1].value)[:5], 'generation_mw':number(row[2].value),
               'shed_mw':number(row[4].value), 'demand_mw':number(row[6].value)}
              for row in workbook['P4'].iter_rows(min_row=9, max_row=33, max_col=7)]
    routes = []
    if 'GenLog' in workbook.sheetnames:
        log = workbook['GenLog']
        for col in range(2, log.max_column + 1):
            name = str(log.cell(11, col).value or '')
            if any(token in name.lower() for token in ['hvdc','tripura','adani','nepal']):
                kwh = number(log.cell(38, col).value)
                routes.append({'name':name, 'energy_mwh':round(kwh / 1000, 6) if kwh is not None else None})
    report = {'date':day.isoformat(), 'report_id':item['id'], 'published_date':item['published_date'],
              'water_level':extract_water_level(workbook),
              'regions':[{'name':str(sheet.cell(row,2).value).replace(' Zone',''),
                          'shed_mw':number(sheet.cell(60+row-17,7).value)}
                         for row in range(17,26) if sheet.cell(row,2).value],
              'operations':extract_event_log(sheet),
              'source_url':item['download_url'], 'fuels':fuels, 'hourly':hourly, 'imports':routes,
              'summary':{'energy_mwh':energy(sheet['K8'].value), 'unserved_mwh':energy(sheet['K9'].value),
                         'import_mwh':energy(sheet['K26'].value), 'temperature_c':temperature(sheet['K11'].value)}}
    notes = ['National totals extracted separately; full workbook parser rejected this report.']
    generated = report['summary']['energy_mwh']
    if generated is None:
        raise ValueError('Missing national daily energy total')
    if all(value is not None for value in fuels.values()) and abs(sum(fuels.values())-generated) >= .1:
        notes.append('Reported fuel subtotals do not reconcile to national energy; fuel mix omitted.')
        report['fuels'] = {fuel:None for fuel in fuels}
    imported = report['summary']['import_mwh']
    if imported is not None and all(route['energy_mwh'] is not None for route in routes) and abs(sum(route['energy_mwh'] for route in routes)-imported) >= .1:
        notes.append('Reported import route subtotals do not reconcile; route breakdown omitted.')
        report['imports'] = []
    report['quality_notes'] = notes
    report['_trends_only'] = True
    return report
