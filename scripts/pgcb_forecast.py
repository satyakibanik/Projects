"""Dated PGCB scenarios and generation-side actuals, used only for comparison."""
import datetime as dt
import io
import math
import re


def number(value):
    return float(value) if type(value) in (float, int) and math.isfinite(value) and value >= 0 else None


def scenario_from_sheet(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    for index, row in enumerate(rows):
        if len(row) < 10 or not str(row[4]).strip().startswith('Forecast Scenario'):
            continue
        match = re.search(r'(\d{2}-\d{2}-\d{4})', str(row[4]))
        try: day = dt.datetime.strptime(match[1], '%d-%m-%Y').date().isoformat() if match else None
        except ValueError: day = None
        hourly = []
        header = next((i for i in range(index+1, min(index+4, len(rows)))
                       if len(rows[i]) >= 10 and rows[i][4] == 'Hour' and rows[i][9] == 'Demand'), None)
        if header is not None:
            for values in rows[header+1:]:
                clock = values[4]
                if clock is None: break
                clock = clock.strftime('%H:%M') if hasattr(clock, 'strftime') else str(clock)[:5]
                hourly.append(dict(time=clock, generation_mw=number(values[5]),
                                   shed_mw=number(values[7]), demand_mw=number(values[9])))
        return day, hourly
    return None, []


def read_workbook(body, item):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=True)
    try:
        day, hourly = scenario_from_sheet(book['Forecast'])
        scenario = dict(date=day, hourly=hourly, report_id=str(item['id']),
                        published_date=item['published_date'], source_url=item['download_url'])
        actual_day, actuals = None, []
        if 'P4' in book:
            sheet = book['P4']
            raw = sheet['C7'].value
            try:
                actual_day = raw.date().isoformat() if isinstance(raw, dt.datetime) else raw.isoformat() if isinstance(raw, dt.date) else dt.datetime.strptime(str(raw), '%d-%m-%Y').date().isoformat()
            except (ValueError, TypeError): pass
            for values in sheet.iter_rows(min_row=9, max_row=40, values_only=True):
                clock = values[1]
                clock = clock.strftime('%H:%M') if hasattr(clock, 'strftime') else str(clock)[:5]
                if re.fullmatch(r'(?:[01]\d|2[0-3]):00', clock):
                    actuals.append(dict(time=clock, demand_mw=number(values[6])))
        return scenario, actual_day, actuals
    finally:
        book.close()


def collect_comparisons(reports, today, refresh=False):
    """Refresh the newest scenarios after long model fits; never use them as inputs."""
    scenarios, actuals, sources = {}, {}, {}
    for report in sorted(reports.values(), key=lambda r: (r.get('published_date',''), str(r.get('report_id','')))):
        forecast = report.get('forecast', {})
        day = forecast.get('scenario_date')
        if day:
            scenarios[day] = dict(date=day, hourly=forecast.get('demand_hourly', []),
                published_date=report.get('published_date'), source_url=report.get('source_url'))
        if report.get('date') and report.get('hourly'):
            actuals[report['date']] = [dict(time=r['time'], demand_mw=r.get('demand_mw')) for r in report['hourly']]
            sources[report['date']] = dict(basis='Workbook P4 · generation side', source_url=report.get('source_url'))
    if refresh:
        import server
        try:
            listing, _ = server.fetch_listing_page(1)
            start = (dt.date.fromisoformat(today)-dt.timedelta(days=1)).isoformat()
            for item in sorted(listing, key=lambda r: (r['published_date'], int(r['id']))):
                if not start <= item['published_date'] <= today or str(item['id']) in reports: continue
                try:
                    scenario, day, rows = read_workbook(server.fetch_bytes(item['download_url'], limit=server.MAX_REPORT_BYTES), item)
                    if scenario['date'] and scenario['hourly']:
                        scenarios[scenario['date']] = scenario
                        print(f"PGCB scenario refreshed: {scenario['date']} from report {item['id']}", flush=True)
                    if day and rows:
                        actuals[day] = rows
                        sources[day] = dict(basis='Workbook P4 · generation side', source_url=item['download_url'])
                except Exception as error:
                    print(f"PGCB report {item['id']} unavailable: {type(error).__name__}: {error}", flush=True)
        except Exception as error:
            print(f'PGCB listing refresh unavailable: {type(error).__name__}: {error}', flush=True)
    return scenarios, actuals, sources
