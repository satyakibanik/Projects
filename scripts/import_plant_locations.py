#!/usr/bin/env python3
"""Build the static plant-location catalog from the user's reviewed matching workbook."""
import argparse
import datetime
import hashlib
import json
import math
import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]


def in_ring(lon, lat, ring):
    inside = False
    for a, b in zip(ring, ring[1:] + ring[:1]):
        cross = (lat - a[1]) * (b[0] - a[0]) - (lon - a[0]) * (b[1] - a[1])
        if abs(cross) < 1e-10 and min(a[0], b[0]) - 1e-10 <= lon <= max(a[0], b[0]) + 1e-10 and min(a[1], b[1]) - 1e-10 <= lat <= max(a[1], b[1]) + 1e-10:
            return True
        if (b[1] > lat) != (a[1] > lat) and lon < (a[0] - b[0]) * (lat - b[1]) / (a[1] - b[1]) + b[0]:
            inside = not inside
    return inside


def inside_country(lon, lat, geometry):
    polygons = [geometry['coordinates']] if geometry['type'] == 'Polygon' else geometry['coordinates']
    return any(in_ring(lon, lat, p[0]) and not any(in_ring(lon, lat, hole) for hole in p[1:]) for p in polygons)


def import_locations(workbook, osm_file, boundary_file, reviewed_on):
    book = openpyxl.load_workbook(workbook, data_only=True)
    sheet = book['Plant matches']
    header = {cell.value: cell.column for row in sheet for cell in row if cell.value == 'Match status'}
    if not header:
        raise ValueError('Missing Match status header')
    header_row = next(cell.row for row in sheet for cell in row if cell.value == 'Match status')
    cols = {cell.value: cell.column for cell in sheet[header_row] if cell.value}
    osm = {e['type'] + '/' + str(e['id']): e for e in json.loads(osm_file.read_text())['elements']}
    geometry = json.loads(boundary_file.read_text())['geometry']
    plants, excluded, seen = [], [], set()
    for r in range(header_row + 1, sheet.max_row + 1):
        value = lambda label: sheet.cell(r, cols[label]).value
        name, status = value('Dashboard plant'), str(value('Match status') or '').strip()
        if not name:
            continue
        normalized = re.sub('[^a-z0-9]', '', name.lower())
        if normalized in seen:
            raise ValueError(f'Duplicate plant key: {name}')
        seen.add(normalized)
        if status.lower() == 'not a plant':
            excluded.append({'name': name, 'reason': 'Not a power plant'})
            continue
        osm_id = value('OSM ID')
        lat, lon = value('Latitude (WGS84)'), value('Longitude (WGS84)')
        if status in ('Matched site', 'Shared complex'):
            source = 'OSM complex center' if status == 'Shared complex' else 'OSM site center'
        elif status.lower() == 'use the candidate':
            center = osm[osm_id].get('center', osm[osm_id])
            lat, lon = center['lat'], center['lon']
            source = 'User-approved OSM candidate'
        else:
            numbers = re.findall(r'\d+\.\d+', status)
            if len(numbers) != 2:
                raise ValueError(f'Unresolved status for {name}: {status}')
            lat, lon = map(float, numbers)
            source, osm_id = 'User-supplied coordinate', None
        if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in (lat, lon)):
            raise ValueError(f'Missing coordinates: {name}')
        if not inside_country(lon, lat, geometry):
            raise ValueError(f'Coordinate outside Bangladesh: {name} ({lat}, {lon})')
        plants.append(dict(name=name, latitude=lat, longitude=lon, source=source, osm_id=osm_id, approximate='?' in status))
    return dict(reviewed_on=reviewed_on, source_workbook=workbook.name,
                source_sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
                coordinate_reference='WGS84', plants=plants, excluded=excluded,
                attribution='Plant locations derived from OpenStreetMap contributors (ODbL) and user-reviewed coordinates.')


def main():
    folder = ROOT / 'Power Plant Coordinates'
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=pathlib.Path, default=folder / 'Power Plant Coordinate Matches.xlsx')
    parser.add_argument('--osm', type=pathlib.Path, default=folder / 'export.json')
    parser.add_argument('--reviewed-on', default=datetime.date.today().isoformat())
    parser.add_argument('--output', type=pathlib.Path, default=ROOT / 'begd/plant-locations.json')
    args = parser.parse_args()
    result = import_locations(args.workbook, args.osm, ROOT / 'begd/bangladesh-boundary.geojson', args.reviewed_on)
    args.output.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':')) + '\n')
    print(f"Wrote {len(result['plants'])} reviewed locations; excluded {len(result['excluded'])} non-plant rows.")


if __name__ == '__main__':
    main()
