"""Offline regression tests; upstream ERP is never contacted."""
import datetime as dt
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import server
import import_report
import build_pages_data as build


def workbook():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in ('P1', 'P2', 'P3', 'P4', 'Forecast', 'GenLog', 'Voltage', 'EWIC'):
        book.create_sheet(name)
    p = book['P1']
    p['D6'] = dt.datetime(2026, 10, 14)
    for row in range(17, 27):
        p.cell(row, 2, f'Region {row} Zone')
        for col in range(5, 14):
            p.cell(row, col, 0)
    for cell in ('K8', 'K9', 'K11', 'K26'):
        p[cell] = 0
    p['K8'] = 2.4
    p['E26'] = 2.4
    p['M17'] = 2.4
    p['E17'] = 2.4
    book['P4']['G9'] = 0
    p2 = book['P2']
    p2['J11'] = 'Normal'
    for col, value in {3: 'Test Plant', 5: 'Public', 6: '2 x 100', 7: 180, 8: 100, 9: 2400000}.items():
        p2.cell(11, col, value)
    p2.cell(12, 2, 'Dhaka Area Total')
    f = book['Forecast']
    f['C6'] = '15-10-2026'
    f['B11'] = 'Test Plant'
    f['C11'] = 'Gas'
    g = book['GenLog']
    g['B11'] = 'Test Plant (Gas)'
    g['B12'] = '180 MW'
    g['B38'] = 2400000
    for hour in range(24):
        g.cell(13 + hour, 1, f'{hour:02}:00')
        g.cell(13 + hour, 2, 100)
    return book


def parse(book):
    output = io.BytesIO()
    book.save(output)
    return import_report.parse_report(output.getvalue(), 'fixture.xlsx')


class WorkbookTests(unittest.TestCase):
    def test_units_capacity_and_reconciliation(self):
        report = parse(workbook())
        self.assertEqual(report['date'], '2026-10-14')
        self.assertEqual(report['forecast']['date'], '2026-10-15')
        self.assertEqual(report['summary']['energy_mwh'], 2400)
        self.assertEqual(report['plants'][0]['region'], 'Dhaka')
        plant = report['plant_units'][0]
        self.assertEqual(plant['installed_capacity_mw'], 200)
        self.assertEqual(plant['fuel'], 'Gas')
        self.assertEqual(plant['operating_hours'], 24)
        self.assertEqual(plant['capacity_factor_pct'], 50)
        self.assertEqual(plant['puf_pct'], 50)

    def test_rejects_inconsistent_energy(self):
        book = workbook()
        book['P2']['I11'] = 100
        with self.assertRaisesRegex(AssertionError, 'Plant energy'):
            parse(book)

    def test_string_operating_date(self):
        book = workbook()
        book['P1']['D6'] = '14-10-2026'
        self.assertEqual(parse(book)['date'], '2026-10-14')

    def test_missing_sheet_rejected(self):
        book = workbook()
        del book['GenLog']
        with self.assertRaises(KeyError):
            parse(book)


class ServerTests(unittest.TestCase):
    def test_numbers(self):
        for text, expected in [(' 1,250.5 ', 1250.5), ('-4', -4), ('—', None), ('-', None), ('bad', None), ('', None)]:
            with self.subTest(text=text):
                self.assertEqual(server.parse_number(text), expected)

    def test_table_nested_text(self):
        parser = server.DataTableParser()
        parser.feed('<table><tr><td>A &amp; <b>B</b></td><td>3</td></tr></table>')
        self.assertEqual(parser.rows, [['A & B', '3']])

    def test_listing_and_pagination(self):
        body = b'''<table><tr><td>1</td><td>Daily report</td><td>15-10-2026</td><td><a href="/web/files/download?location=erp/web/report_docs/123.xlsx">Download</a></td></tr></table>Showing 1 of 25'''
        rows, pages = server.parse_listing(body)
        self.assertEqual(pages, 25)
        self.assertEqual(rows[0]['id'], '123')
        self.assertEqual(rows[0]['published_date'], '2026-10-15')
        self.assertTrue(rows[0]['download_url'].startswith('https://erp.powergrid.gov.bd/'))

    def test_empty_listing(self):
        with self.assertRaises(ValueError):
            server.parse_listing(b'<html>unavailable</html>')
        self.assertEqual(server.parse_listing(b'', require_count=False), ([], None))

    def test_external_urls_rejected_before_network(self):
        for url in ('http://erp.powergrid.gov.bd/a', 'https://example.com/a', 'file:///etc/passwd'):
            with self.subTest(url=url), self.assertRaises(ValueError):
                server.fetch_bytes(url)

    def test_current_table_normalizes_and_sorts(self):
        body = b'<tr><td>15-10-2026</td><td>9:30:00</td><td>1</td></tr><tr><td>15-10-2026</td><td>8:00</td><td>2</td></tr>'
        with patch.object(server, 'fetch_bytes', return_value=body):
            rows = server.parse_current_table(minimum_cells=3, url='https://example.test')
        self.assertEqual([r[1] for r in rows], ['08:00', '09:30'])

    def test_current_table_rejects_missing_readings(self):
        with patch.object(server, 'fetch_bytes', return_value=b'<tr><td>Header</td></tr>'):
            with self.assertRaises(LookupError):
                server.parse_current_table(minimum_cells=3, url='https://example.test')

    def test_current_day_cache_hit(self):
        with patch.object(server, 'current_day_cache', {'today_date': '2026-10-15'}), patch.object(server, 'current_day_cache_time', 100), patch.object(server.time, 'time', return_value=101), patch.object(server, 'fetch_current_day') as fetch:
            self.assertEqual(server.get_current_day()['today_date'], '2026-10-15')
            fetch.assert_not_called()


class BuildTests(unittest.TestCase):
    def test_archive_empty(self):
        with patch.object(server, 'fetch_listing_page', return_value=([], 0)):
            with self.assertRaises(RuntimeError):
                build.available_reports()

    def test_bundle_escapes_script_and_selects_operating_day(self):
        records = [dict(id='2', name='</script>&', published_date='2026-10-15'), dict(id='1', name='Older', published_date='2026-10-16')]
        def parsed(item):
            return dict(report_id=item['id'], published_date=item['published_date'], date='2026-10-14', name=item['name'])
        with tempfile.TemporaryDirectory() as directory, patch.object(build, 'available_reports', return_value=records), patch.object(build, 'parse_listing_item', side_effect=parsed), patch.object(server, 'get_highest_generation_day', side_effect=LookupError), patch.object(server, 'get_current_day', return_value={'today_date': '2026-10-15'}):
            output = Path(directory)
            build.write_data_bundle(output)
            text = (output / 'live-data.js').read_text()
            self.assertNotIn('</script>', text)
            bundle = json.loads(text.removeprefix('window.BEGD_STATIC_DATA=').strip().removesuffix(';'))
            self.assertEqual(bundle['reports_by_operating_date']['2026-10-14'], '2')
            self.assertEqual(len(bundle['reports']), 2)
            index = (output / 'index.html').read_text()
            self.assertLess(index.index('live-data.js'), index.index('static-api.js'))
            self.assertLess(index.index('static-api.js'), index.index('src="app.js'))

    def test_bundle_survives_upstream_outage(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(build, 'available_reports', side_effect=RuntimeError), patch.object(server, 'get_highest_generation_day', side_effect=LookupError), patch.object(server, 'get_current_day', side_effect=TimeoutError):
            build.write_data_bundle(Path(directory))
            text = (Path(directory) / 'live-data.js').read_text()
            bundle = json.loads(text.removeprefix('window.BEGD_STATIC_DATA=').strip().removesuffix(';'))
            self.assertEqual(bundle['reports'], [])
            self.assertEqual(bundle['failures'], ['archive: RuntimeError'])
            self.assertEqual(bundle['current_day']['demand_hourly'], [])


if __name__ == '__main__':
    unittest.main()
