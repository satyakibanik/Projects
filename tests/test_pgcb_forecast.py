"""PGCB comparison data must stay dated, separate from model inputs, and basis-correct."""
import datetime as dt
import io
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import pgcb_forecast as pgcb
import erp_history as erp
import build_pages_data as build
import daily_forecast as daily
from test_daily_forecast import ready


def workbook():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Forecast'
    sheet['C6'] = '25-10-2026'  # Plant availability differs from the hourly scenario.
    sheet['E10'] = 'Forecast Scenario (24-10-2026)'
    sheet['E11'], sheet['J11'] = 'Hour', 'Demand'
    for h in range(24):
        sheet.cell(12+h,5,f'{h:02}:00')
        sheet.cell(12+h,6,100+h)
        sheet.cell(12+h,8,0)
        sheet.cell(12+h,10,101+h)
    actual = book.create_sheet('P4')
    actual['C7'] = dt.datetime(2026,10,23)
    for h in range(24):
        actual.cell(9+h,2,f'{h:02}:00')
        actual.cell(9+h,3,90+h)
        actual.cell(9+h,7,91+h)
    output=io.BytesIO(); book.save(output)
    return output.getvalue()


class PGCBTests(unittest.TestCase):
    def test_reader_uses_printed_hourly_date_and_retains_separate_p4_actual_date(self):
        scenario, day, actuals = pgcb.read_workbook(workbook(), dict(id='5174',published_date='2026-10-24',download_url='https://example.test/24.xlsx'))
        self.assertEqual(scenario['date'],'2026-10-24')
        self.assertEqual(len(scenario['hourly']),24)
        self.assertEqual(day,'2026-10-23')
        self.assertEqual(actuals[0],dict(time='00:00',demand_mw=91))
        self.assertEqual(scenario['hourly'][0]['shed_mw'],0)

    def test_new_report_after_long_fit_supplies_scenario_and_matching_basis_actuals(self):
        item=dict(id='5174',published_date='2026-10-24',download_url='https://example.test/24.xlsx')
        with patch('server.fetch_listing_page',return_value=([item],1)), patch('server.fetch_bytes',return_value=workbook()):
            scenarios, actuals, sources = pgcb.collect_comparisons({},'2026-10-24',True)
        self.assertEqual(scenarios['2026-10-24']['report_id'],'5174')
        self.assertEqual(len(actuals['2026-10-23']),24)
        self.assertEqual(sources['2026-10-23']['basis'],'Workbook P4 · generation side')
        self.assertNotIn('2026-10-24',actuals)

    def test_listing_outage_keeps_existing_matching_scenarios_and_actuals(self):
        reports={'23':dict(date='2026-10-22',published_date='2026-10-23',hourly=[dict(time='00:00',demand_mw=10)],forecast=dict(scenario_date='2026-10-23',demand_hourly=[dict(time='00:00',demand_mw=12)]))}
        with patch('server.fetch_listing_page',side_effect=TimeoutError):
            scenarios, actuals, _ = pgcb.collect_comparisons(reports,'2026-10-24',True)
        self.assertEqual(set(scenarios),{'2026-10-23'})
        self.assertEqual(actuals['2026-10-22'][0]['demand_mw'],10)

    def test_generation_side_erp_reader_requires_the_right_basis_and_preserves_query(self):
        body=b'<tr><th>Date</th><th>Time</th><th>Demand (MW) Generation Side</th><th>Generation (MW)</th><th>Loadshed</th><th>Remarks</th></tr><tr><td>24-10-2026</td><td>00:00</td><td>105</td><td>100</td><td>5</td><td></td></tr>'
        rows,_=erp.parse_table(body,'generation_demand','2026-10-24')
        self.assertEqual(rows[0]['demand_mw'],105)
        with self.assertRaises(ValueError): erp.parse_table(body.replace(b'Generation Side',b'Substation Side'),'generation_demand')
        with tempfile.TemporaryDirectory() as directory,patch.object(erp,'download',return_value=body) as fetch:
            erp.load_range('generation_demand','2026-10-24','2026-10-24',Path(directory))
        self.assertIn('?d_gen=1&search=24-10-2026&page=1',fetch.call_args.args[0])

    def test_publication_refreshes_pgcb_after_model_preparation_without_using_it_in_fitting(self):
        events=[]
        def prepare(*args):
            events.append('fit')
            self.assertNotIn('pgcb_forecast',args[1])
            return {'2026-10-24':ready('2026-10-24')}
        def listing(*args):
            events.append('scenario refresh')
            return [dict(id='5174',published_date='2026-10-24',download_url='https://example.test/24.xlsx')],1
        current=dict(today_date='2026-10-24')
        with tempfile.TemporaryDirectory() as directory,patch.object(daily,'prepare_forecasts',side_effect=prepare),patch('server.fetch_listing_page',side_effect=listing),patch('server.fetch_bytes',return_value=workbook()),patch.object(erp,'load_range',return_value=dict(days={})):
            build.attach_forecasts([],{},current,Path(directory),refresh_pgcb=True)
        self.assertEqual(events,['fit','scenario refresh'])
        self.assertEqual(current['pgcb_forecast']['date'],'2026-10-24')
        self.assertEqual(current['forecast_history']['2026-10-23']['generation_side_demand_hourly'][0]['demand_mw'],91)
        self.assertEqual(current['generation_side_demand_hourly'],[])
