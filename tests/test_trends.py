"""Historical energy aggregation must preserve source dates and missingness."""
import datetime as dt
import sys
from pathlib import Path
import unittest
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from trends import build_trends, daily_summary


class TrendsTests(unittest.TestCase):
    def test_regional_shedding_preserves_snapshots_zero_and_missing(self):
        row = daily_summary({'date':'2026-10-29', 'hourly':[{'shed_mw':100}],
                             'regions':[{'name':'Dhaka','shed_mw':0}, {'name':'Khulna','shed_mw':20},
                                        {'name':'Sylhet','shed_mw':None}, {'name':'Rangpur','shed_mw':-2}]})
        self.assertEqual(row['regional_shed_mw'], {'Dhaka':0, 'Khulna':20, 'Sylhet':None, 'Rangpur':None})
        self.assertEqual(row['shed_peak_mw'],100)
        self.assertEqual(daily_summary({'date':'2026-10-29'})['regional_shed_mw'],{})

    def test_daily_energy_uses_published_totals_not_hourly_sums(self):
        report = {'date':'2026-10-18', 'summary':{'energy_mwh':2400,'unserved_mwh':0,'import_mwh':120},
                  'hourly':[{'generation_mw':99999,'shed_mw':0},{'shed_mw':35},{'shed_mw':None}],
                  'en_curve':{'date':'2026-10-18','totals':{'total_requirement_mwh':2500}},
                  'fuels':{'Gas':2200,'Imports':120,'Solar':0},
                  'imports':[{'name':'Nepal','energy_mwh':0}]}
        row = daily_summary(report)
        self.assertEqual(row['generation_mwh'],2400)
        self.assertEqual(row['demand_mwh'],2500)
        self.assertEqual(row['shed_peak_mw'],35)
        self.assertEqual(row['unserved_mwh'],0)
        self.assertEqual(row['routes']['Nepal'],0)

    def test_requirement_falls_back_to_reported_p1_energy_with_provenance(self):
        row = daily_summary({'date':'2026-01-01', 'summary':{'energy_mwh':2400, 'unserved_mwh':125}})
        self.assertEqual(row['demand_mwh'],2525)
        self.assertEqual(row['demand_basis'],'P1 generation + reported energy not served')
        row = daily_summary({'date':'2026-01-01', 'summary':{'energy_mwh':2400}})
        self.assertIsNone(row['demand_mwh'])

    def test_missing_and_invalid_values_are_not_zero(self):
        row = daily_summary({'date':'2026-01-01','summary':{'energy_mwh':float('nan'), 'import_mwh':-1},
                             'hourly':[{'shed_mw':None},{'shed_mw':True},{'shed_mw':-5}],
                             'fuels':{'Gas':None,'Coal':float('inf')}})
        for field in ['generation_mwh','demand_mwh','shed_peak_mw','imports_mwh','unserved_mwh']:
            self.assertIsNone(row[field])
        self.assertIsNone(row['fuels']['Coal'])

    def test_temperature_preserves_signed_zero_and_missing_readings(self):
        for value in [0, -5, 33.7]:
            row = daily_summary({'date':'2026-01-01','summary':{'temperature_c':value}})
            self.assertEqual(row['temperature_c'],value)
        for value in [None, True, float('nan'), float('inf'), '33.7']:
            row = daily_summary({'date':'2026-01-01','summary':{'temperature_c':value}})
            self.assertIsNone(row['temperature_c'])

    def test_demand_curve_must_match_operating_day(self):
        row=daily_summary({'date':'2026-01-01','en_curve':{'date':'2026-01-02','totals':{'total_requirement_mwh':500}}})
        self.assertIsNone(row['demand_mwh'])

    def test_duplicate_operating_days_have_archive_precedence(self):
        reports=[{'date':'2026-01-01','published_date':published,'report_id':key,'summary':{'energy_mwh':value}}
                 for published,key,value in [('2026-01-03','3',300),('2026-01-02','1',100),('2026-01-02','2',200)]]
        result=build_trends(reports,'now')
        self.assertEqual(len(result['days']),1)
        self.assertEqual(result['days'][0]['generation_mwh'],200)
        self.assertEqual(result['days'][0]['report_id'],'2')

    def test_365_days_are_inclusive_and_old_reports_are_excluded(self):
        end=dt.date(2026,10,18)
        reports=[{'date':(end-dt.timedelta(days=i)).isoformat(),'report_id':str(i+1)} for i in range(400)]
        result=build_trends(reports,'now')
        self.assertEqual(result['end_date'],'2026-10-18')
        self.assertEqual(len(result['days']),365)
        self.assertEqual(result['days'][0]['date'],'2025-10-19')
        self.assertEqual(build_trends([],'now')['days'],[])


class NationalRecoveryTests(unittest.TestCase):
    def test_ancillary_failure_does_not_discard_national_totals(self):
        import io
        import openpyxl
        from trends import extract_trend_workbook
        book=openpyxl.Workbook();book.active.title='P1';book.create_sheet('P4');book.create_sheet('En-Curve')
        p=book['P1'];p['D6']=dt.datetime(2026,9,11);p['K8']=240;p['K9']=5;p['K26']=0;p['K11']=33.7
        p['E26']=240
        for col in [6,7,8,9,10,11,12]:p.cell(26,col,0)
        book['P4']['E9']=125
        p['B17']='Dhaka Zone';p['G60']=0
        p['B18']='Chattogram Zone';p['G61']=25
        book['En-Curve']['B2']='bad ancillary date'
        stream=io.BytesIO();book.save(stream)
        report=extract_trend_workbook(stream.getvalue(),{'id':'5129','published_date':'2026-09-12','download_url':'https://erp.powergrid.gov.bd/report'})
        row=daily_summary(report)
        self.assertEqual(row['regional_shed_mw'],{'Dhaka':0,'Chattogram':25})
        self.assertEqual(row['date'],'2026-09-11');self.assertEqual(row['generation_mwh'],240000)
        self.assertEqual(row['demand_mwh'],245000);self.assertEqual(row['shed_peak_mw'],125)
        self.assertEqual(row['temperature_c'],33.7)
        self.assertTrue(report['_trends_only']);self.assertTrue(row['quality_notes'])
        p['E26']=200;stream=io.BytesIO();book.save(stream)
        report=extract_trend_workbook(stream.getvalue(),{'id':'5129','published_date':'2026-09-12','download_url':'https://erp.powergrid.gov.bd/report'})
        self.assertEqual(report['summary']['energy_mwh'],240000)
        self.assertIsNone(report['fuels']['Gas'])


class BengalGridTests(unittest.TestCase):
    def test_display_filter_preserves_raw_source_and_aggregates_observations(self):
        import tempfile
        from build_bengalgrid_trends import aggregate
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'source.csv'
            source='datetime,Generation_MW\n2026-01-01T00:00:00,1000\n2026-01-01T00:30:00,3000\n2026-01-01T01:00:00,64526500\n2026-01-03T00:00:00,4000\n'
            path.write_text(source);data=aggregate(path)
            self.assertEqual(path.read_text(),source)
            self.assertEqual(data['days'][0]['mean_mw'],2000)
            self.assertEqual(data['days'][0]['observations'],2)
            self.assertEqual(data['days'][0]['excluded'],1)
            self.assertEqual(len(data['days']),2)
            self.assertEqual(data['months'][0]['median_mw'],3000)
            self.assertEqual(data['months'][0]['days'],2)
            self.assertEqual(len(data['excluded_records']),1)

    def test_bundled_historical_data_matches_paper_exclusions(self):
        import json
        data=json.loads((Path(__file__).resolve().parents[1]/'begd/bengalgrid-trends.json').read_text())
        self.assertEqual(len(data['excluded_records']),7)
        self.assertEqual(data['start_date'],'2015-04-17');self.assertEqual(data['end_date'],'2026-10-07')
        self.assertTrue(all(756<=row['mean_mw']<=17468 for row in data['days']))
        self.assertEqual(data['license'],'CC BY 4.0')


class FallbackBuildTests(unittest.TestCase):
    def setUp(self):
        from unittest.mock import patch
        costs = patch('bpdb_costs.collect', return_value={'days':[], 'errors':[]})
        costs.start()
        self.addCleanup(costs.stop)

    def test_recovered_totals_are_published_without_exposing_partial_full_reports(self):
        import tempfile
        import json
        from unittest.mock import patch
        import build_pages_data as build
        import server
        item={'id':'91234','name':'Daily Report','published_date':'2026-09-12','download_url':'https://erp.powergrid.gov.bd/report'}
        partial={'date':'2026-09-11','report_id':'91234','published_date':'2026-09-12',
                 '_trends_only':True,'summary':{'energy_mwh':240000},'quality_notes':['National totals only']}
        with tempfile.TemporaryDirectory() as directory, patch.object(build,'ROOT',Path(directory)), patch.object(build,'available_reports',return_value=[item]), patch.object(build,'parse_listing_item',side_effect=AssertionError('Plant mismatch')), patch.object(server,'fetch_bytes',return_value=b'workbook'), patch.object(build,'extract_trend_workbook',return_value=partial), patch.object(server,'get_highest_generation_day',side_effect=LookupError), patch.object(server,'get_current_day',return_value={}):
            output=Path(directory)/'output';build.write_data_bundle(output)
            trends=json.loads((output/'trends-data.json').read_text())
            self.assertEqual(trends['days'][0]['generation_mwh'],240000)
            self.assertTrue(trends['days'][0]['quality_notes'])
            bundle=json.loads((output/'live-data.js').read_text().split('=',1)[1].rstrip(';\n'))
            self.assertEqual(bundle['reports'],[])
            self.assertEqual(bundle['reports_by_id'],{})
