"""Event extraction, report precedence, fallback and cache-backfill contracts."""
import datetime as dt
import io
from pathlib import Path
import sys
import unittest
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from operational_events import extract_event_log, build_event_log, event_cache_current
from trends import extract_trend_workbook
from test_scripts import workbook, parse


class EventTests(unittest.TestCase):
    def test_moving_headers_preserve_first_entry_and_wrapped_description(self):
        for header in [81, 82, 83, 84]:
            sheet = openpyxl.Workbook().active
            sheet.cell(header, 4, 'Description of Event')
            sheet.cell(header+1, 2, '08:01')
            sheet.cell(header+1, 4, 'Plant Forced Shutdown Due to')
            sheet.cell(header+2, 4, 'complete wrapped source description')
            sheet.cell(header+3, 3, '09:20')
            sheet.cell(header+3, 4, 'Line is restored.')
            result = extract_event_log(sheet)
            self.assertEqual(len(result['events']), 2)
            self.assertEqual(result['events'][0]['description'], 'Plant Forced Shutdown Due to complete wrapped source description')
            self.assertEqual(result['events'][0]['source_rows'], [header+1, header+2])
            self.assertEqual(result['events'][1]['outage_time'], '—')
            self.assertEqual(result['events'][1]['restoration_time'], '09:20')

    def test_missing_header_and_empty_log_are_distinct(self):
        sheet = openpyxl.Workbook().active
        self.assertIsNone(extract_event_log(sheet)['events'])
        sheet['D83'] = 'Description of Event'
        self.assertEqual(extract_event_log(sheet)['events'], [])

    def test_zero_clock_and_untimed_description_are_not_discarded(self):
        sheet = openpyxl.Workbook().active
        sheet['D83'] = 'Description of Event'
        sheet['B84'] = dt.time(0, 0)
        sheet['D84'] = 'Unit Shutdown'
        sheet['C86'] = 'Sub-Divisional Engineer'
        sheet['D88'] = '.'
        result = extract_event_log(sheet)['events']
        self.assertEqual(result[0]['outage_time'], '00:00')
        self.assertEqual(result[1]['description'], '.')

    def test_full_and_fallback_parsers_share_complete_event_extraction(self):
        book = workbook()
        sheet = book['P1']
        sheet['D83'] = 'Description of Event'
        sheet['B84'] = '08:02'
        sheet['D84'] = 'Scheduled S/D Due to'
        sheet['D85'] = 'annual maintenance'
        report = parse(book)
        output = io.BytesIO()
        book.save(output)
        fallback = extract_trend_workbook(output.getvalue(), dict(id='2',published_date='2026-10-15',download_url='https://example.test/report'))
        self.assertEqual(report['operations']['events'], fallback['operations']['events'])
        self.assertEqual(report['operations']['event_log_version'], 1)
        self.assertEqual(report['operations']['events'][0]['source_rows'], [84, 85])

    def test_bundle_uses_trends_end_date_window_and_report_precedence(self):
        def report(day, rid, published, events):
            return dict(date=day,report_id=rid,published_date=published,operations=dict(event_log_version=1,events=events))
        reports=[report('2026-10-29','3','2026-10-30',[]),report('2026-10-29','4','2026-10-31',[dict(description='wrong revision')]),report('2025-10-29','1','2025-10-30',[]),report('2025-10-30','2','2025-10-31',None),report('2026-10-30','5','2026-10-31',[])]
        result=build_event_log(reports,'snapshot','2026-10-29')
        self.assertEqual([d['date'] for d in result['days']], ['2025-10-30','2026-10-29'])
        self.assertEqual(result['days'][-1]['report_id'],'3')
        self.assertEqual(result['days'][-1]['events'],[])
        self.assertIsNone(result['days'][0]['events'])
        self.assertEqual(build_event_log([], 'snapshot', None)['days'], [])

    def test_old_cache_is_unavailable_until_refreshed(self):
        old=dict(date='2026-10-29',operations=dict(events=[dict(description='partial')]))
        self.assertIsNone(build_event_log([old],'now','2026-10-29')['days'][0]['events'])
        latest=dt.date(2026,10,30)
        self.assertFalse(event_cache_current(old,'2026-10-29',latest))
        self.assertTrue(event_cache_current(old,'2025-01-01',latest))
        old['operations']['event_log_version']=1
        self.assertTrue(event_cache_current(old,'2026-10-29',latest))

if __name__=='__main__':unittest.main()
